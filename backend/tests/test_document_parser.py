"""EXT-4 文档解析注册表的回归测试。

要守住的三件事：
    1. **单一真相**：后缀集合、富文本集合都从注册表派生，工具/索引/预览不再各存一份；
    2. **新格式一处生效**：注册一个解析器后，`supported_suffixes()` 与 `is_structured()` 同步变化；
    3. **降级不炸**：未知后缀按文本兜底、结构化解析失败回落文本读取。

Excel 用例真实调用 openpyxl（已装），顺带验证 sheet 名与截断标记。
"""

# 让类型注解延迟求值（PEP 563）
from __future__ import annotations

# 标准库：临时目录与路径
from pathlib import Path

# 测试框架
import pytest

# 项目内：被测试的注册表
from app.services import document_parser as dp


def test_builtin_suffixes_registered() -> None:
    """内置解析器要覆盖 pdf/docx/文本/xlsx，并正确区分结构化与纯文本。"""
    supported = dp.supported_suffixes()
    assert {".pdf", ".docx", ".txt", ".md", ".markdown", ".xlsx"} <= supported

    structured = dp.structured_suffixes()
    assert {".pdf", ".docx", ".xlsx"} <= structured
    assert ".txt" not in structured
    assert ".md" not in structured


def test_is_structured_case_insensitive(tmp_path: Path) -> None:
    """后缀判断不区分大小写（Windows 上用户经常写成 .PDF）。"""
    assert dp.is_structured("D:/a/b.PDF") is True
    assert dp.is_structured(tmp_path / "x.DocX") is True
    assert dp.is_structured("D:/a/b.txt") is False
    assert dp.is_structured("D:/a/b.unknown") is False


def test_parse_text_file(tmp_path: Path) -> None:
    """纯文本按编码读出来。"""
    path = tmp_path / "note.md"
    path.write_text("# 标题\n正文", encoding="utf-8")
    assert dp.parse_document(path) == "# 标题\n正文"


def test_parse_unknown_suffix_falls_back_to_text(tmp_path: Path) -> None:
    """未注册后缀（如 .log）按文本兜底，而不是报"不支持"。"""
    path = tmp_path / "app.log"
    path.write_text("2026-10-20 boot ok", encoding="utf-8")
    assert "boot ok" in dp.parse_document(path)


def test_gbk_file_is_decoded(tmp_path: Path) -> None:
    """GBK 编码的中文文件要能读（Windows 上很常见）。"""
    path = tmp_path / "gbk.txt"
    path.write_bytes("中文内容".encode("gbk"))
    assert dp.parse_document(path) == "中文内容"


def test_parse_xlsx_keeps_sheet_names(tmp_path: Path) -> None:
    """Excel 逐 sheet 读值，并保留 [Sheet 名] 以便检索定位。"""
    from openpyxl import Workbook

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "预算"
    sheet.append(["项目", "金额"])
    sheet.append(["服务器", 1200])
    second = workbook.create_sheet("备注")
    second.append(["说明", "按月结算"])
    path = tmp_path / "book.xlsx"
    workbook.save(path)

    text = dp.parse_document(path)
    assert "[Sheet 预算]" in text
    assert "服务器 | 1200" in text
    assert "[Sheet 备注]" in text
    assert "按月结算" in text


def test_xlsx_row_cap(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    """行数超上限时要显式标注截断，而不是静默丢数据。"""
    from openpyxl import Workbook

    monkeypatch.setattr(dp, "MAX_XLSX_ROWS", 3)
    workbook = Workbook()
    sheet = workbook.active
    for index in range(6):
        sheet.append([f"r{index}"])
    path = tmp_path / "big.xlsx"
    workbook.save(path)

    text = dp.parse_document(path)
    assert "已截断" in text
    assert "r4" not in text


def test_register_new_parser_takes_effect(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    """注册一个新后缀后，派生集合与解析行为同时生效（这就是"一处生效"的验收点）。"""
    monkeypatch.setitem(dp.PARSERS, ".demo", dp.ParserSpec(
        suffix=".demo", parse=lambda p: f"DEMO:{p.name}", structured=True, label="演示格式",
    ))
    assert ".demo" in dp.supported_suffixes()
    assert ".demo" in dp.structured_suffixes()

    path = tmp_path / "a.demo"
    path.write_text("x", encoding="utf-8")
    assert dp.parse_document(path) == "DEMO:a.demo"


def test_structured_parse_failure_falls_back(tmp_path: Path) -> None:
    """结构化解析器抛错时回落文本读取（例如损坏的 docx），不把异常抛给调用方。"""
    def broken(_path: Path) -> str:
        """模拟解析器内部报错。"""
        raise RuntimeError("corrupted")

    path = tmp_path / "fake.docx"
    path.write_text("其实是纯文本", encoding="utf-8")

    original = dp.PARSERS[".docx"]
    try:
        dp.PARSERS[".docx"] = dp.ParserSpec(suffix=".docx", parse=broken, structured=True, label="坏解析器")
        assert "其实是纯文本" in dp.parse_document(path)
    finally:
        dp.PARSERS[".docx"] = original
